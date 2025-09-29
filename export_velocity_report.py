# export_velocity_report.py
from __future__ import annotations
import math
import os
import numpy as np
from typing import Optional, Tuple
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter


def _to_nhw(x):
    x = np.asarray(x)
    if x.ndim == 4 and x.shape[1] == 1:  # (N,1,H,W)
        return x[:, 0, ...]
    if x.ndim == 4 and x.shape[-1] == 1:  # (N,H,W,1)
        return x[..., 0]
    return x  # (N,H,W)


def _r2_score(y_true, y_pred, mask=None):
    if mask is not None:
        y_true = y_true[mask]
        y_pred = y_pred[mask]
    ss_res = np.sum((y_true - y_pred) ** 2)
    ss_tot = np.sum((y_true - np.mean(y_true)) ** 2)
    return 1.0 - (ss_res / ss_tot if ss_tot > 0 else np.nan)


def _pearsonr(y_true, y_pred, mask=None):
    if mask is not None:
        y_true = y_true[mask]
        y_pred = y_pred[mask]
    y_true = y_true.ravel()
    y_pred = y_pred.ravel()
    if y_true.size < 2:
        return np.nan
    yt = y_true - y_true.mean()
    yp = y_pred - y_pred.mean()
    denom = (np.linalg.norm(yt) * np.linalg.norm(yp))
    return float(yt.dot(yp) / denom) if denom > 0 else np.nan


def _basic_metrics(y_true, y_pred, mask=None):
    if mask is not None:
        y_true = y_true[mask]
        y_pred = y_pred[mask]
    diff = (y_pred - y_true).astype(np.float64)
    mae  = np.mean(np.abs(diff))
    mse  = np.mean(diff ** 2)
    rmse = math.sqrt(mse)
    r2   = _r2_score(y_true, y_pred)
    r    = _pearsonr(y_true, y_pred)
    maxae = np.max(np.abs(diff)) if diff.size else np.nan
    # MAPE (safe): only on non-zero ground truth
    nz = np.abs(y_true) > 1e-8
    mape = np.mean(np.abs(diff[nz] / y_true[nz])) if nz.any() else np.nan
    return dict(MAE=mae, MSE=mse, RMSE=rmse, R2=r2, PearsonR=r, MaxAE=maxae, MAPE=mape)


def _write_table(ws, start_row, start_col, arr, title=None, as_heatmap=True):
    """Write a 2D array as a cell grid and optionally add a heatmap rule."""
    r0, c0 = start_row, start_col
    if title:
        ws.cell(r0, c0, title).font = Font(bold=True)
        r0 += 1
    H, W = arr.shape
    # write numbers
    for i in range(H):
        for j in range(W):
            ws.cell(r0 + i, c0 + j, float(arr[i, j]))
    # auto column widths (cap to keep sheet readable)
    max_w = min(W, 64)
    for j in range(max_w):
        ws.column_dimensions[get_column_letter(c0 + j)].width = 10

    # add 3-color heatmap
    if as_heatmap:
        max_row = r0 + H - 1
        max_col = c0 + W - 1
        ref = f"{get_column_letter(c0)}{r0}:{get_column_letter(max_col)}{max_row}"
        rule = ColorScaleRule(
            start_type="percentile", start_value=5, start_color="63BE7B",  # green
            mid_type="percentile",   mid_value=50, mid_color="FFEB84",     # yellow
            end_type="percentile",   end_value=95, end_color="F8696B",     # red
        )
        ws.conditional_formatting.add(ref, rule)

    return r0 + H, c0 + W  # next free row/col


def _write_kv(ws, row, col, mapping, title=None):
    if title:
        ws.cell(row, col, title).font = Font(bold=True, size=12)
        row += 1
    for k, v in mapping.items():
        ws.cell(row, col, k)
        ws.cell(row, col + 1, float(v) if isinstance(v, (int, float, np.floating)) else v)
        row += 1
    return row


def _make_mask(den_slice: Optional[np.ndarray]) -> Optional[np.ndarray]:
    """If a density mask is passed, treat zeros as 'solid' to exclude."""
    if den_slice is None:
        return None
    den_slice = np.asarray(den_slice)
    if den_slice.ndim == 3:
        den_slice = den_slice[..., 0]
    return den_slice > 0.0


def save_velocity_report(
    vx_pred,
    vy_pred,
    vx_true=None,
    vy_true=None,
    den: Optional[np.ndarray] = None,
    out_dir: str = "Results/prediction_report",
    out_xlsx: str = "predictions.xlsx",
    max_sheets: Optional[int] = None,
    title: str = "Model Predictions (compact report)",
):
    """
    Build a SINGLE Excel file with:
      - 'Summary' sheet: overall and per-image metrics
      - One sheet per image: metrics + cell-by-cell tables (heatmap) for vx/vy (pred, truth, error)

    Inputs may be shaped (N,H,W), (N,1,H,W) or (N,H,W,1). 'true' arrays optional for prediction-only runs.
    """
    vx_pred = _to_nhw(vx_pred)
    vy_pred = _to_nhw(vy_pred)
    vx_true = _to_nhw(vx_true) if vx_true is not None else None
    vy_true = _to_nhw(vy_true) if vy_true is not None else None
    den     = _to_nhw(den)     if den is not None else None

    N = vx_pred.shape[0]
    H, W = vx_pred.shape[1:3]
    if max_sheets is None:
        max_sheets = N
    K = min(max_sheets, N)

    os.makedirs(out_dir, exist_ok=True)
    xlsx_path = os.path.join(out_dir, out_xlsx)

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"

    # Header
    ws_sum["A1"] = title
    ws_sum["A1"].font = Font(bold=True, size=14)
    _write_kv(
        ws_sum, 3, 1,
        {"Num Images": K, "Height": H, "Width": W, "Masked by den?": den is not None},
        title="Dataset Info"
    )

    # Per-image metrics table
    start_row = 8
    headers = ["idx", "channel", "MAE", "MSE", "RMSE", "R2", "PearsonR", "MaxAE", "MAPE"]
    for j, h in enumerate(headers, start=1):
        ws_sum.cell(start_row, j, h).font = Font(bold=True)
    row = start_row + 1

    # Collect overall metrics (averaged over images)
    overall_accum = {("vx", k): [] for k in headers[2:]}
    overall_accum.update({("vy", k): [] for k in headers[2:]})

    for i in range(K):
        mask = _make_mask(den[i]) if den is not None else None

        # vx metrics
        vx_metrics = _basic_metrics(vx_true[i], vx_pred[i], mask) if vx_true is not None else {k: np.nan for k in headers[2:]}
        ws_sum.cell(row, 1, i)
        ws_sum.cell(row, 2, "vx")
        for j, key in enumerate(headers[2:], start=3):
            val = vx_metrics.get(key, np.nan)
            ws_sum.cell(row, j, float(val) if val is not None else np.nan)
            overall_accum[("vx", key)].append(val)
        row += 1

        # vy metrics
        vy_metrics = _basic_metrics(vy_true[i], vy_pred[i], mask) if vy_true is not None else {k: np.nan for k in headers[2:]}
        ws_sum.cell(row, 1, i)
        ws_sum.cell(row, 2, "vy")
        for j, key in enumerate(headers[2:], start=3):
            val = vy_metrics.get(key, np.nan)
            ws_sum.cell(row, j, float(val) if val is not None else np.nan)
            overall_accum[("vy", key)].append(val)
        row += 1

    # Overall averages
    row += 1
    ws_sum.cell(row, 1, "OVERALL AVERAGES").font = Font(bold=True)
    row += 1
    ws_sum.cell(row, 1, "channel").font = Font(bold=True)
    for j, key in enumerate(headers[2:], start=2):
        ws_sum.cell(row, j, key).font = Font(bold=True)
    row += 1
    for ch in ("vx", "vy"):
        ws_sum.cell(row, 1, ch)
        for j, key in enumerate(headers[2:], start=2):
            vals = np.array(overall_accum[(ch, key)], dtype=float)
            with np.errstate(invalid="ignore"):
                ws_sum.cell(row, j, float(np.nanmean(vals)) if vals.size else np.nan)
        row += 1

    # Per-image sheets
    for i in range(K):
        ws = wb.create_sheet(title=f"img_{i}")
        ws["A1"] = f"Image {i}"
        ws["A1"].font = Font(bold=True, size=13)

        # Metrics box
        mask = _make_mask(den[i]) if den is not None else None
        vx_metrics = _basic_metrics(vx_true[i], vx_pred[i], mask) if vx_true is not None else {}
        vy_metrics = _basic_metrics(vy_true[i], vy_pred[i], mask) if vy_true is not None else {}
        row = 3
        row = _write_kv(ws, row, 1, {"Channel": "vx"}, title="Metrics")
        row = _write_kv(ws, row, 1, vx_metrics or {"(no ground truth)": ""})
        row += 1
        row = _write_kv(ws, row, 1, {"Channel": "vy"})
        row = _write_kv(ws, row, 1, vy_metrics or {"(no ground truth)": ""})

        # Cell-by-cell tables with heatmaps
        # Layout block: vx at cols 6.., vy below it
        start_col = 6
        r = 2
        r, _ = _write_table(ws, r, start_col, vx_pred[i], title="vx_pred")
        r, _ = _write_table(ws, r + 1, start_col, (vx_true[i] if vx_true is not None else np.zeros_like(vx_pred[i])), title="vx_true")
        r, _ = _write_table(ws, r + 1, start_col, (vx_pred[i] - vx_true[i]) if vx_true is not None else np.zeros_like(vx_pred[i]), title="vx_error")

        r = r + 3
        r, _ = _write_table(ws, r, start_col, vy_pred[i], title="vy_pred")
        r, _ = _write_table(ws, r + 1, start_col, (vy_true[i] if vy_true is not None else np.zeros_like(vy_pred[i])), title="vy_true")
        _ , _ = _write_table(ws, r + 1, start_col, (vy_pred[i] - vy_true[i]) if vy_true is not None else np.zeros_like(vy_pred[i]), title="vy_error")

        # Tidy captions
        for cell in ("A1", "F2"):
            if ws[cell].value:
                ws[cell].alignment = Alignment(vertical="center")

    # Save one compact file
    wb.save(xlsx_path)
    return xlsx_path
